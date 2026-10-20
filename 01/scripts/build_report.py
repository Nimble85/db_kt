#!/usr/bin/env python3
"""Складає звіт до практичної роботи 1 у форматі DOCX."""
import os
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Cm, Pt

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SHOTS = os.path.join(BASE, "screenshots")
LOGS = os.path.join(BASE, "logs")
OUT_DIR = os.path.join(os.path.dirname(BASE), "src", "ouput")
OUT = os.path.join(OUT_DIR, "Звіт_ПР1.docx")

TITLE = {
    "zvo": "Чорноморський національний університет імені Петра Могили",
    "faculty": "Факультет комп'ютерних наук",
    "chair": "Кафедра інтелектуальних інформаційних систем",
    "course": "Технології баз даних та знань",
    "student": "Климошенко Сергій",
    "group": "K26-1м",
    "specialty": "Компʼютерні науки",
    "teacher": "________________________",
    "city": "Миколаїв",
    "year": "2026",
}
THEME = ("Робота зі СКБД SQLite. Створення БД. Створення таблиці. "
         "Імпорт-експорт, резервне копіювання та відновлення")

FIG_W_PER_CHAR = 0.175
FIG_W_MAX = 16.5
FIG_W_MIN = 5.0

CONTENT = [
    ("h1", "Тема та мета роботи"),
    ("p", "Тема: " + THEME + "."),
    ("p", "Мета: навчитись працювати з СКБД SQLite, створювати БД, робити імпорт-експорт "
          "БД, резервне копіювання та відновлення."),
    ("p", "Роботу виконано в середовищі macOS за допомогою консольного клієнта sqlite3 "
          "версії 3.51.0. Методичні вказівки орієнтовані на Windows (командний рядок cmd, "
          "SQLite 3.21.0), тому відповідні кроки виконано в оболонці macOS (запрошення «$»). "
          "Усі зображення є знімками реального сеансу роботи з інтерпретатором. Слід "
          "зауважити, що у цій версії SQLite повідомлення про помилки мають вигляд "
          "«Runtime error: ... (19)» замість «Error: ...», а запрошення має вигляд "
          "«sqlite>» (без пробілу). Для імпорту CSV використано синтаксис "
          "«.import --csv --skip 1», який пропускає рядок заголовка."),

    ("h1", "Хід виконання основного завдання"),
    ("h2", "1. Підготовка середовища та запуск sqlite3"),
    ("p", "Перевіряємо наявність клієнта sqlite3 у системі та його версію (рис. 1). "
          "На macOS установлення не потрібне: sqlite3 входить до складу системи."),
    ("fig", "main_fig01", "Перевірка наявності та версії sqlite3"),
    ("p", "Переходимо до робочої директорії (рис. 2)."),
    ("fig", "main_fig02", "Перехід до потрібної директорії"),
    ("p", "Запускаємо sqlite3 без параметрів та виводимо довідку командою .help "
          "(рис. 3, наведено початок виводу). Без назви БД робота ведеться в тимчасовій "
          "БД в пам'яті. Сеанс завершується командою .exit."),
    ("fig", "main_fig03", "Запуск sqlite3 та команда .help (початок виводу)"),

    ("h2", "2. Створення бази даних"),
    ("p", "Для створення БД достатньо вказати її назву під час запуску sqlite3. "
          "Якщо файл існує, робота продовжується в контексті цієї БД, інакше вона "
          "створюється автоматично. Виконуємо sqlite3 myFirstDB.db; після цього "
          "з'являється запрошення sqlite> для введення наступної команди (рис. 4)."),
    ("fig", "main_fig04", "Створення нової БД. Запрошення sqlite>"),
    ("p", "Перевіряємо список баз даних командою .databases (рис. 5). Файл БД на "
          "диску з'являється після першого запису даних."),
    ("fig", "main_fig06", "Перевірка списку БД. Команда .databases"),

    ("h2", "3. Створення таблиці та наповнення даними"),
    ("p", "Створюємо таблицю students та додаємо до неї три рядки (рис. 6)."),
    ("fig", "main_fig07", "Створення таблиці «студенти» та заповнення її даними"),
    ("p", "Налаштовуємо відображення імен колонок та табличний режим виводу "
          "(рис. 7) і перевіряємо вміст таблиці (рис. 8)."),
    ("fig", "main_fig08", "Налаштування виведення шапки таблиці"),
    ("fig", "main_fig09", "Виведення вмісту таблиці на екран"),
    ("p", "Схему поточної БД виводимо командою .schema (рис. 9). Перелік об'єктів БД "
          "також можна отримати із системної таблиці sqlite_master (рис. 10)."),
    ("fig", "main_fig10", "Виведення схеми (структури) таблиці"),
    ("fig", "main_fig11", "Звернення до системної таблиці sqlite_master"),

    ("h2", "4. Резервне копіювання"),
    ("p", "Командою .dump створюємо резервну копію БД у вигляді SQL-скрипта; для цього "
          "виведення перенаправляємо у файл командою .output і потім повертаємо на "
          "екран (рис. 11). Вміст створеного файлу myDump.sql наведено на рис. 12."),
    ("fig", "main_fig12", "Створення резервної копії БД"),
    ("fig", "main_fig13", "Вміст файлу резервної копії БД"),
    ("p", "Дамп також можна створити з командного рядка ОС (рис. 13). Порівняння "
          "показує, що обидва файли ідентичні."),
    ("fig", "main_fig14", "Створення резервної копії БД із командного рядка ОС"),

    ("h2", "5. Відновлення БД із резервної копії"),
    ("p", "Видаляємо файл myFirstDB.db і відновлюємо БД, виконавши скрипт резервної "
          "копії (рис. 14)."),
    ("fig", "main_fig15", "Відновлення БД із резервної копії"),
    ("p", "Підключаємось до відновленої БД (рис. 15), виводимо список таблиць (рис. 16) "
          "та вибираємо дані з таблиці students (рис. 17). Таблиця та дані "
          "відновлені повністю."),
    ("fig", "main_fig16", "Підключення до БД"),
    ("fig", "main_fig17", "Виведення списку таблиць"),
    ("fig", "main_fig18", "Вибірка даних із таблиці «студенти»"),

    ("h2", "6. Експорт та імпорт даних (CSV)"),
    ("p", "Експортуємо дані таблиці students до зовнішнього файлу stud.csv: вмикаємо "
          "режим csv, перенаправляємо вивід у файл, виконуємо вибірку та переглядаємо "
          "файл (рис. 18). Рядок заголовка потрапив до файлу, оскільки було ввімкнено "
          ".headers on."),
    ("fig", "main_fig19", "Експорт даних до зовнішнього csv файлу"),
    ("p", "Видаляємо дані з таблиці та переконуємось, що вона порожня (рис. 19)."),
    ("fig", "main_fig20", "Видалення даних із таблиці «студенти»"),
    ("p", "Завантажуємо дані з файлу до таблиці (рис. 20) і перевіряємо її вміст "
          "(рис. 21)."),
    ("fig", "main_fig21", "Завантаження даних із csv файлу"),
    ("fig", "main_fig22", "Перевірка вмісту таблиці"),

    ("h1", "Індивідуальне завдання"),
    ("p", "Умова: у окремій БД створити 2 таблиці із обмеженнями цілісності первинного "
          "ключа для довільної предметної області; заповнити таблиці не менше ніж "
          "10 записами в кожній; навчитись видаляти записи, робити експорт таблиці "
          "до зовнішнього файлу, резервне копіювання тощо."),
    ("p", "Предметна область — «бібліотека». БД library.db містить таблицю authors "
          "(автори) та таблицю books (книги); у кожній первинний ключ id "
          "INTEGER PRIMARY KEY AUTOINCREMENT, а books.author_id пов'язано із "
          "authors.id зовнішнім ключем. Обмеження задано безпосередньо у CREATE TABLE, "
          "оскільки ALTER TABLE у SQLite не дозволяє додавати їх до наявних стовпців."),
    ("p", "Створюємо БД library.db (рис. 22) та вмикаємо контроль зовнішніх ключів: "
          "за замовчуванням він вимкнений і вмикається в кожному з'єднанні (рис. 23)."),
    ("fig", "ind_ind01", "Створення БД library.db"),
    ("fig", "ind_ind02", "Увімкнення контролю зовнішніх ключів"),
    ("p", "Створюємо таблиці authors та books (рис. 24–25) і переглядаємо схему "
          "(рис. 26)."),
    ("fig", "ind_ind03", "Створення таблиці «автори»"),
    ("fig", "ind_ind04", "Створення таблиці «книги»"),
    ("fig", "ind_ind05", "Список таблиць та схема БД"),
    ("p", "Заповнюємо таблицю authors десятьма записами (рис. 27) та таблицю books "
          "дванадцятьма записами (рис. 28)."),
    ("fig", "ind_ind06", "Наповнення таблиці «автори» (10 записів)"),
    ("fig", "ind_ind07", "Наповнення таблиці «книги» (12 записів)"),
    ("p", "Перевіряємо роботу первинного ключа: спроба вставити запис із наявним "
          "значенням id відхиляється (рис. 29). Спроба вставити книгу з неіснуючим "
          "автором порушує зовнішній ключ (рис. 30)."),
    ("fig", "ind_ind08", "Порушення унікальності первинного ключа"),
    ("fig", "ind_ind09", "Порушення зовнішнього ключа"),
    ("p", "Видалення записів: книгу з id = 12 та автора з id = 10 (без книг) "
          "видалено успішно, а автора з id = 1 видалити неможливо, бо на нього "
          "посилаються книги (рис. 31)."),
    ("fig", "ind_ind10", "Видалення записів"),
    ("p", "Експортуємо обидві таблиці до файлів authors.csv та books.csv (рис. 32)."),
    ("fig", "ind_ind11", "Експорт таблиць до csv файлів"),
    ("p", "Створюємо резервні копії: SQL-скрипт командою .dump та двійкову копію БД "
          "командою .backup (рис. 33)."),
    ("fig", "ind_ind12", "Резервне копіювання БД"),
    ("p", "Імітуємо втрату даних — видаляємо всі записи (рис. 34) — і відновлюємо "
          "їх імпортом із csv (рис. 35)."),
    ("fig", "ind_ind13", "Видалення всіх записів із таблиць"),
    ("fig", "ind_ind14", "Імпорт даних із csv файлів та перевірка зв'язку таблиць"),
    ("p", "Також перевіряємо відновлення нової БД із SQL-дампу та цілісність "
          "двійкової резервної копії (рис. 36)."),
    ("fig", "ind_ind15", "Відновлення БД із дампу та перевірка резервної копії"),

    ("h1", "Висновки"),
    ("p", "У ході роботи засвоєно основи роботи з СКБД SQLite: створено БД і таблиці "
          "у консольному клієнті sqlite3, виконано вставку та вибірку даних, "
          "переглянуто схему БД і системну таблицю sqlite_master."),
    ("p", "Опановано резервне копіювання (.dump, .backup, дамп з командного рядка ОС), "
          "відновлення БД із резервної копії, експорт даних до CSV та імпорт із CSV "
          "(.import --csv --skip 1)."),
    ("p", "В індивідуальному завданні створено БД «бібліотека» з двома таблицями, "
          "первинними ключами та зовнішнім ключем, по 10 і більше записів у "
          "кожній. Експериментально підтверджено, що первинний ключ не допускає "
          "дублікатів, а контроль зовнішніх ключів потрібно вмикати окремо "
          "(pragma foreign_keys = on) у кожному з'єднанні. Відмінності версії "
          "SQLite 3.51.0 від методички (формат повідомлень про помилки, "
          "синтаксис .import) на результат не впливають."),

    ("h1", "Список використаних джерел"),
    ("num", "Kreibich Jay. Using SQLite. Small. Fast. Reliable. Choose Any Three. — "
            "O'Reilly Media, 2010. — 530 p."),
    ("num", "Allen Grant, Owens Mike. The Definitive Guide to SQLite. — Apress, 2011. — 368 p."),
    ("num", "Sibsankar Haldar. SQLite Database System Design and Implementation "
            "(Second Edition, Version 2). — O'Reilly, 2016. — 264 p."),
    ("num", "What Is SQLite? [Електронний ресурс]. — Режим доступу: "
            "https://www.sqlite.org/index.html. — Загол. з екрану."),
    ("num", "Кренке Д. Теория и практика построения баз данных. 8-е изд. — СПб. : "
            "Питер, 2003. — 800 с."),
    ("num", "Дейт К. Дж. Введение в системы баз данных. 7-е издание : пер. с англ. — "
            "М. : Издательский дом «Вильямс», 2001. — 1072 с."),
    ("num", "Грофф Дж., Вайнберг П. SQL : Полное руководство : пер. с англ. — 2-е изд. — "
            "К. : Издательская группа BHV, 2001. — 816 с."),
    ("num", "Фісун М. Т., Ніколенко С. Г. Створення та ведення баз даних засобами "
            "мови Jet SQL : методичні вказівки. — Миколаїв : вид-во ЧДУ ім. Петра "
            "Могили, 2009. — 83 с."),
    ("num", "Пасічник В. В., Резніченко В. А. Організація баз даних та знань. — К. : "
            "Видавнича група BHV, 2006. — 384 с."),
    ("num", "SQLite Documentation. Command Line Shell For SQLite [Електронний ресурс]. — "
            "Режим доступу: https://www.sqlite.org/cli.html"),
    ("num", "SQLite Documentation. Query Language Understood by SQLite [Електронний "
            "ресурс]. — Режим доступу: https://www.sqlite.org/lang.html"),
    ("num", "SQLite Documentation. SQLite Foreign Key Support [Електронний ресурс]. — "
            "Режим доступу: https://www.sqlite.org/foreignkeys.html"),
]


def setup_styles(doc):
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(14)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts")
    if rf is not None:
        rf.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", "Times New Roman")
    pf = st.paragraph_format
    pf.line_spacing = 1.5
    pf.first_line_indent = Cm(1.25)
    pf.space_after = Pt(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for sec in doc.sections:
        sec.top_margin = Cm(2)
        sec.bottom_margin = Cm(2)
        sec.left_margin = Cm(3)
        sec.right_margin = Cm(1.5)


def para(doc, text="", *, indent=Cm(1.25), align=WD_ALIGN_PARAGRAPH.JUSTIFY,
         bold=False, size=14, space_before=0, space_after=0, spacing=1.5):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.first_line_indent = indent
    pf.alignment = align
    pf.line_spacing = spacing
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = "Times New Roman"
    return p


def title_page(doc, t):
    c = WD_ALIGN_PARAGRAPH.CENTER
    para(doc, "МІНІСТЕРСТВО ОСВІТИ І НАУКИ УКРАЇНИ", indent=0, align=c, spacing=1.0)
    para(doc, t["zvo"].upper(), indent=0, align=c, bold=True, spacing=1.0)
    para(doc, t["faculty"], indent=0, align=c, spacing=1.0)
    para(doc, t["chair"], indent=0, align=c, spacing=1.0, space_after=120)
    para(doc, "ЗВІТ", indent=0, align=c, bold=True, size=16, spacing=1.0, space_before=60)
    para(doc, "до практичної роботи № 1", indent=0, align=c, spacing=1.0)
    para(doc, "з дисципліни «%s»" % t["course"], indent=0, align=c, spacing=1.0)
    para(doc, "на тему: «%s»" % THEME, indent=0, align=c, spacing=1.0, space_after=160)
    r = WD_ALIGN_PARAGRAPH.RIGHT
    para(doc, "Виконав(ла): студент(ка) групи %s" % t["group"], indent=0, align=r,
         spacing=1.0, space_before=80)
    para(doc, "спеціальності %s" % t["specialty"], indent=0, align=r, spacing=1.0)
    para(doc, t["student"], indent=0, align=r, spacing=1.0)
    para(doc, "Перевірив(ла): %s" % t["teacher"], indent=0, align=r, spacing=1.0, space_before=40)
    para(doc, "Оцінка: ______________", indent=0, align=r, spacing=1.0, space_before=40)
    para(doc, "%s – %s" % (t["city"], t["year"]), indent=0, align=c, spacing=1.0, space_before=200)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def fig_width_cm(key):
    with open(os.path.join(LOGS, key + ".txt")) as f:
        chars = max(len(l) for l in f.read().split("\n"))
    return min(FIG_W_MAX, max(FIG_W_MIN, chars * FIG_W_PER_CHAR))


def add_figure(doc, key, caption, number):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = 0
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(os.path.join(SHOTS, key + ".png"), width=Cm(fig_width_cm(key)))
    cp = doc.add_paragraph()
    cp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.first_line_indent = 0
    cp.paragraph_format.space_after = Pt(10)
    cp.paragraph_format.line_spacing = 1.0
    r1 = cp.add_run("Рис. %d. " % number)
    r1.bold = True
    r1.font.size = Pt(12)
    r1.font.name = "Times New Roman"
    r2 = cp.add_run(caption)
    r2.font.size = Pt(12)
    r2.font.name = "Times New Roman"


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    doc = Document()
    setup_styles(doc)
    title_page(doc, TITLE)
    fig_no = li_no = 0
    for item in CONTENT:
        kind = item[0]
        if kind == "h1":
            para(doc, item[1], indent=0, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True,
                 space_before=12, space_after=8).paragraph_format.keep_with_next = True
        elif kind == "h2":
            para(doc, item[1], bold=True, space_before=10, space_after=4).paragraph_format.keep_with_next = True
        elif kind == "p":
            para(doc, item[1])
        elif kind == "num":
            li_no += 1
            para(doc, "%d. %s" % (li_no, item[1]))
        elif kind == "fig":
            fig_no += 1
            add_figure(doc, item[1], item[2], fig_no)
    doc.save(OUT)
    print("Звіт збережено: %s; рисунків: %d" % (OUT, fig_no))


if __name__ == "__main__":
    main()
