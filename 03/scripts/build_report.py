#!/usr/bin/env python3
"""Складає звіт до практичної роботи 3 у форматі DOCX."""
import os
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SHOTS = os.path.join(BASE, "screenshots")
LOGS = os.path.join(BASE, "logs")
OUT = os.path.join(os.path.dirname(BASE), "src", "ouput", "Звіт_ПР3.docx")

TITLE = {
    "zvo": "Чорноморський національний університет імені Петра Могили",
    "faculty": "Факультет комп'ютерних наук",
    "chair": "Кафедра інтелектуальних інформаційних систем",
    "course": "Технології баз даних та знань",
    "student": "Климошенко Сергій",
    "group": "K26-1м",
    "specialty": "Компʼютерні науки",
    "teacher": "________________________",
    "city": "Миколаїв",
    "year": "2026",
}

FIG_W_PER_CHAR = 0.175   # см на символ моноширинного тексту
FIG_W_MAX = 16.5
FIG_W_MIN = 5.0

# --------------------------------------------------------------------------
# Зміст звіту: ("h1"|"h2"|"p"|"li"|"fig", ...)
# --------------------------------------------------------------------------
CONTENT = [
    ("h1", "Тема та мета роботи"),
    ("p", "Тема: робота зі СКБД SQLite. Виконання запитів на вибірку даних."),
    ("p", "Мета: навчитись засобами СКБД SQLite виконувати запити на вибірку даних."),
    ("p", "Роботу виконано в середовищі macOS за допомогою консольного клієнта "
          "sqlite3 версії 3.51.0. Усі наведені нижче зображення є знімками реального "
          "сеансу роботи з інтерпретатором. Відмінності від методичних вказівок, "
          "зумовлені новішою версією СКБД (у методичці — 3.21.0): режим column "
          "автоматично підбирає ширину стовпців, тому назви на кшталт «Chicken/Fowl» "
          "не обрізаються; результати довгих запитів на знімках скорочено позначкою "
          "«[... вивід скорочено ...]». Ілюстративні схеми методички (загальна форма "
          "оператора select, конвеєр обробки, етапи виконання) подано в тексті, а не "
          "знімками консолі."),

    ("h1", "Хід виконання основного завдання"),

    ("h2", "1. Тестова база даних"),
    ("p", "Надану студентом БД foods.db містить крім чотирьох таблиць предметної "
          "області (episodes, foods, food_types, foods_episodes) ще допоміжну порожню "
          "таблицю contacts, яка в завданні не використовується. Із чотирьох потрібних "
          "таблиць командою .dump сформовано дамп foods.sql, який відновлюється в нову "
          "БД food.db (рис. {main_fig01}); далі переглядаємо перелік таблиць (рис. {main_fig02})."),
    ("fig", "main_fig01", "Відновлення БД із дампу foods.sql та запуск sqlite3"),
    ("fig", "main_fig02", "Перегляд переліку таблиць БД food"),
    ("p", "БД містить назви страв, що згадуються в епізодах серіалу «Seinfeld». "
          "Foods — головна таблиця: кожен запис є окремою стравою, назва зберігається в "
          "полі name, а поле type_id посилається на таблицю food_types (класифікація: "
          "фаст-фуд, напої тощо). Таблиця foods_episodes пов'язує страви з епізодами "
          "(зв'язок «багато до багатьох»). Схему БД (даталогічну модель) наведено "
          "командою .schema (рис. {main_fig03})."),
    ("fig", "main_fig03", "Даталогічна модель БД «Food» (команда .schema)"),

    ("h2", "2. Загальна форма оператора select"),
    ("p", "Найповніша форма оператора вибірки в SQLite має вигляд: select [distinct] "
          "heading from tables where predicate group by columns having predicate "
          "order by columns limit count,offset. Майже всі фрази необов'язкові. "
          "Оператор можна уявити як конвеєр: фраза from поєднує таблиці (join) в одне "
          "відношення R1; where відбирає записи (обмеження) та утворює R2; group by "
          "групує записи й застосовує агрегатні функції; having фільтрує групи; select "
          "виконує проекцію полів; order by упорядковує результат; limit обмежує "
          "кількість рядків."),

    ("h2", "3. Прості вибірки"),
    ("p", "Виконуємо просту вибірку зі списку типів страв (рис. {main_fig07}) та "
          "проекцію — лише назви типів (рис. {main_fig08}). Режим column та виведення "
          "заголовків увімкнено командами .mode column та .headers on."),
    ("fig", "main_fig07", "Вибірка зі списку типів страв"),
    ("fig", "main_fig08", "Назви типів страв"),
    ("p", "Вибираємо назви всіх серій серіалу (рис. {main_fig09}). Кількість рядків "
          "велика, тому за допомогою limit обмежуємо результат п'ятьма серіями "
          "(рис. {main_fig10}), а також п'ятьма першими стравами (рис. {main_fig11})."),
    ("fig", "main_fig09", "Вибірка назв серій серіалу (вивід скорочено)"),
    ("fig", "main_fig10", "Обмеження результуючого списку фразою limit"),
    ("fig", "main_fig11", "Перші 5 страв зі списку страв"),
    ("p", "Для зручності виводу змінюємо ширину стовпців командою .width (рис. {main_fig12})."),
    ("fig", "main_fig12", "Вказання ширини стовпців при виводі даних"),

    ("h2", "4. Фільтрування даних"),
    ("p", "Список серій другого сезону — рис. {main_fig13}. Умови поєднуються "
          "булівськими операторами: страви 2-го та 3-го типу (рис. {main_fig14}); ту ж "
          "вибірку реалізовано оператором in (рис. {main_fig15})."),
    ("fig", "main_fig13", "Список серій 2-го сезону"),
    ("fig", "main_fig14", "Список страв 2-го та 3-го типу (вивід скорочено)"),
    ("fig", "main_fig15", "Використання оператора in (вивід скорочено)"),
    ("p", "Під час фільтрації за рядковим значенням застосовуємо шаблони оператора like: "
          "страви, назва яких пов'язана з морквою (рис. {main_fig16}), та серії "
          "3-го сезону зі словом «boyfriend» у назві (рис. {main_fig17})."),
    ("fig", "main_fig16", "Страви, назва яких пов'язана з морквою"),
    ("fig", "main_fig17", "Серії 3-го сезону, в назві яких є слово «boyfriend»"),

    ("h2", "5. Групування та агрегатні функції"),
    ("p", "Визначаємо загальну кількість серій (рис. {main_fig18}), кількість серій "
          "2-го сезону (рис. {main_fig19}) та кількість сезонів із застосуванням "
          "count(distinct ...) (рис. {main_fig20}). Зверніть увагу: у таблиці episodes є "
          "два службові записи без номера сезону (season = NULL: id 0 та id 180, "
          "спеціальний випуск), тому count(*) дає 181 запис, а кількість різних сезонів — 9."),
    ("fig", "main_fig18", "Загальна кількість серій у серіалі"),
    ("fig", "main_fig19", "Кількість серій 2-го сезону"),
    ("fig", "main_fig20", "Кількість сезонів у серіалі"),
    ("p", "Речення group by розбиває вибірку на групи: підрахунок серій за сезонами "
          "(рис. {main_fig21}); службові записи з NULL утворюють окрему групу. Результат "
          "агрегації фільтруємо реченням having — сезони, у яких більше 20 серій "
          "(рис. {main_fig22})."),
    ("fig", "main_fig21", "Підрахунок кількості серій за сезонами"),
    ("fig", "main_fig22", "Сезони, в яких більше 20 серій"),

    ("h2", "6. Упорядкування виводу"),
    ("p", "Список типів страв за назвою (рис. {main_fig23}); страви 1-го типу у "
          "зворотному порядку (рис. {main_fig24}); серії за сезоном за зростанням і "
          "назвою за спаданням (рис. {main_fig25}); сезони за кількістю серій за "
          "спаданням — упорядкування за результатом агрегатної функції (рис. {main_fig26})."),
    ("fig", "main_fig23", "Список типів страв, упорядкований за назвою"),
    ("fig", "main_fig24", "Список страв 1-го типу, відсортований за назвою у зворотному порядку (вивід скорочено)"),
    ("fig", "main_fig25", "Список серій, відсортований за сезоном (зростання) та назвою серії (спадання) (вивід скорочено)"),
    ("fig", "main_fig26", "Список сезонів, впорядкований за кількістю серій за спаданням"),

    ("h2", "7. Запити до кількох таблиць"),
    ("p", "Для поєднання таблиць застосовуємо конструкцію join. Список страв із вказанням "
          "їхніх типів — рис. {main_fig27}; салати з категорії овочі — рис. {main_fig28}."),
    ("fig", "main_fig27", "Список страв із вказанням їхніх типів"),
    ("fig", "main_fig28", "Салати із категорії овочі"),
    ("p", "Приклади з поєднанням join та group by: перші 5 епізодів за кількістю "
          "згадувань про їжу (рис. {main_fig29}); перші 3 сезони за кількістю згадувань "
          "різних страв (рис. {main_fig30}). У п'ятому рядку рейтингу епізодів у методичці "
          "стоїть «The Dinner Party», а в нашому виконанні — «The Wink»: обидва мають по "
          "9 згадувань, тож порядок рівних значень не визначено, і він може відрізнятися "
          "залежно від версії СКБД."),
    ("fig", "main_fig29", "Перші 5 епізодів за кількістю згадувань про їжу"),
    ("fig", "main_fig30", "Перші 3 сезони за кількістю згадувань різних страв"),
    ("p", "У скількох серіях кожного сезону згадується морозиво — рис. {main_fig31}; "
          "у яких серіях 5-го сезону згадується випічка — рис. {main_fig32}."),
    ("fig", "main_fig31", "У скількох серіях кожного із сезонів згадується про морозиво"),
    ("fig", "main_fig32", "У яких серіях 5-го сезону згадується випічка"),

    ("h1", "Індивідуальне завдання"),
    ("p", "Для БД «Страви» (окрема копія food_ind.db, відновлена з того самого дампу) "
          "реалізовано сім запитів на вибірку."),
    ("p", "1. Типи страв, що починаються з літери «D» (рис. {ind_ind01}). Оператор like "
          "у SQLite нечутливий до регістру для латинських літер."),
    ("fig", "ind_ind01", "Типи страв, що починаються з літери «D»"),
    ("p", "2. Страви з кодом категорії від 3 до 5 — оператор between (рис. {ind_ind02})."),
    ("fig", "ind_ind02", "Страви із кодом категорії від 3 до 5 (вивід скорочено)"),
    ("p", "3. Серії 9 сезону, в назві яких зустрічається «fin» (рис. {ind_ind03}). Крім "
          "очікуваних «The Finale 1» та «The Finale 2» вибірка містить «The Merv Griffin "
          "Show», оскільки «Griffin» містить підрядок «fin»."),
    ("fig", "ind_ind03", "Серії 9 сезону, в назві яких є «fin»"),
    ("p", "4. Типи страв із кількістю страв у кожному (рис. {ind_ind04}); використано "
          "left join, щоб у результат потрапили й типи без страв (таких немає)."),
    ("fig", "ind_ind04", "Типи страв із кількістю страв, що входять до них"),
    ("p", "5. Перші 3 типи страв за кількістю страв (рис. {ind_ind05})."),
    ("fig", "ind_ind05", "Перші 3 типи страв за кількістю страв"),
    ("p", "6. Страва, що згадується у найбільшій кількості епізодів (рис. {ind_ind06}). "
          "Найближчий конкурент має 4 епізоди, тож лідер єдиний."),
    ("fig", "ind_ind06", "Страва, що згадується у найбільшій кількості епізодів"),
    ("p", "7. Кількість різних типів страв, що згадуються в кожному із сезонів "
          "(рис. {ind_ind07}); порожнє значення сезону відповідає службовим записам "
          "без номера сезону."),
    ("fig", "ind_ind07", "Кількість різних типів страв у кожному із сезонів"),

    ("h1", "Висновки"),
    ("p", "У ході виконання практичної роботи було опановано засоби СКБД SQLite для "
          "виконання запитів на вибірку даних."),
    ("li", "Оператор select працює як конвеєр: from (з'єднання) → where (обмеження) → "
           "group by / having (групування й фільтрація груп) → select (проекція) → "
           "order by → limit; всі фрази, крім select, необов'язкові."),
    ("li", "Фільтрування виконується предикатами з оператором in, like, between та "
           "булівськими операторами; like нечутливий до регістру для ASCII, що слід "
           "враховувати (наприклад, «fin» знаходить і «Griffin»)."),
    ("li", "Агрегатні функції count, count(distinct ...) разом із group by і having "
           "дозволяють підраховувати показники по групах; order by допускає "
           "упорядкування за кількома полями та за результатом агрегації."),
    ("li", "Для запитів до кількох таблиць застосовується join; поєднання join із "
           "group by дозволяє отримувати рейтинги (епізоди, сезони, типи страв)."),
    ("li", "У даних є особливості, які впливають на результат: два епізоди без номера "
           "сезону (NULL) утворюють окрему групу, а за рівних значень агрегата порядок "
           "рядків не визначений."),
    ("p", "Основне та індивідуальне завдання виконано в повному обсязі, усі запити "
          "виконано в реальному sqlite3 версії 3.51.0."),

    ("h1", "Список використаних джерел"),
    ("num", "Кренке Д. Теория и практика построения баз данных. 8-е изд. — СПб. : "
            "Питер, 2003. — 800 с."),
    ("num", "Дейт К. Дж. Введение в системы баз данных. 7-е издание : пер. с англ. — "
            "М. : Издательский дом «Вильямс», 2001. — 1072 с."),
    ("num", "Грофф Дж., Вайнберг П. SQL : Повне керівництво : пер. з англ. — 2-ге вид. — "
            "К. : Видавнича група BHV, 2001. — 816 с."),
    ("num", "Конноллі Т., Бегг К. Базы данных. Проектирование, реализация и "
            "сопровождение. Теория и практика. 3-е издание : пер. с англ. — М. : "
            "Издательский дом «Вильямс», 2003. — 1440 с."),
    ("num", "Фісун М. Т., Ніколенко С. Г. Створення та ведення баз даних засобами мови "
            "Jet SQL : методичні вказівки. — Миколаїв : вид-во ЧДУ ім. Петра Могили, "
            "2009. — 83 с."),
    ("num", "SQLite Documentation. SELECT [Електронний ресурс]. — Режим доступу: "
            "https://www.sqlite.org/lang_select.html"),
    ("num", "SQLite Documentation. Aggregate Functions [Електронний ресурс]. — Режим "
            "доступу: https://www.sqlite.org/lang_aggfunc.html"),
    ("num", "SQLite Documentation. Command Line Shell For SQLite [Електронний ресурс]. — "
            "Режим доступу: https://www.sqlite.org/cli.html"),
]

# Підстановка номерів рисунків: {ключ} -> N
_nums = {it[1]: i + 1 for i, it in enumerate([c for c in CONTENT if c[0] == "fig"])}
CONTENT = [((c[0], c[1].format(**_nums)) + tuple(c[2:])) if c[0] == "p" else c
           for c in CONTENT]


def setup_styles(doc):
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(14)
    pf = st.paragraph_format
    pf.line_spacing = 1.5
    pf.first_line_indent = Cm(1.25)
    pf.space_after = Pt(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for sec in doc.sections:
        sec.top_margin = Cm(2)
        sec.bottom_margin = Cm(2)
        sec.left_margin = Cm(3)
        sec.right_margin = Cm(1.5)


def para(doc, text="", *, indent=Cm(1.25), align=WD_ALIGN_PARAGRAPH.JUSTIFY,
         bold=False, size=14, space_before=0, space_after=0, spacing=1.5):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.first_line_indent = indent
    pf.alignment = align
    pf.line_spacing = spacing
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = "Times New Roman"
    return p


def title_page(doc, t):
    c = WD_ALIGN_PARAGRAPH.CENTER
    para(doc, "МІНІСТЕРСТВО ОСВІТИ І НАУКИ УКРАЇНИ", indent=0, align=c, spacing=1.0)
    para(doc, t["zvo"].upper(), indent=0, align=c, bold=True, spacing=1.0)
    para(doc, t["faculty"], indent=0, align=c, spacing=1.0)
    para(doc, t["chair"], indent=0, align=c, spacing=1.0, space_after=120)
    para(doc, "ЗВІТ", indent=0, align=c, bold=True, size=16, spacing=1.0,
         space_before=60)
    para(doc, "до практичної роботи № 3", indent=0, align=c, size=14, spacing=1.0)
    para(doc, "з дисципліни «%s»" % t["course"], indent=0, align=c, spacing=1.0)
    para(doc, "на тему: «Робота зі СКБД SQLite. Виконання запитів на вибірку даних»",
         indent=0, align=c, spacing=1.0, space_after=160)

    r = WD_ALIGN_PARAGRAPH.RIGHT
    para(doc, "Виконав(ла): студент(ка) групи %s" % t["group"], indent=0, align=r,
         spacing=1.0, space_before=80)
    para(doc, "спеціальності %s" % t["specialty"], indent=0, align=r, spacing=1.0)
    para(doc, t["student"], indent=0, align=r, spacing=1.0)
    para(doc, "Перевірив(ла): %s" % t["teacher"], indent=0, align=r, spacing=1.0,
         space_before=40)
    para(doc, "Оцінка: ______________", indent=0, align=r, spacing=1.0,
         space_before=40)

    para(doc, "%s – %s" % (t["city"], t["year"]), indent=0,
         align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.0, space_before=200)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_figure(doc, key, caption, number):
    png = os.path.join(SHOTS, key + ".png")
    log = os.path.join(LOGS, key + ".txt")
    with open(log) as f:
        chars = max(len(l) for l in f.read().split("\n"))
    width = min(FIG_W_MAX, max(FIG_W_MIN, chars * FIG_W_PER_CHAR))

    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = 0
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.line_spacing = 1.0
    p.add_run().add_picture(png, width=Cm(width))

    cp = doc.add_paragraph()
    cp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.first_line_indent = 0
    cp.paragraph_format.space_after = Pt(10)
    cp.paragraph_format.line_spacing = 1.0
    r1 = cp.add_run("Рис. %d. " % number)
    r1.bold = True
    r1.font.size = Pt(12)
    r1.font.name = "Times New Roman"
    r2 = cp.add_run(caption)
    r2.font.size = Pt(12)
    r2.font.name = "Times New Roman"


def main():
    doc = Document()
    setup_styles(doc)
    title_page(doc, TITLE)

    fig_no = 0
    li_no = 0
    for item in CONTENT:
        kind = item[0]
        if kind == "h1":
            para(doc, item[1], indent=0, align=WD_ALIGN_PARAGRAPH.CENTER,
                 bold=True, space_before=12, space_after=8)
        elif kind == "h2":
            para(doc, item[1], bold=True, space_before=10, space_after=4)
        elif kind == "p":
            para(doc, item[1])
        elif kind == "li":
            p = para(doc, "— " + item[1])
            p.paragraph_format.first_line_indent = Cm(1.25)
        elif kind == "num":
            li_no += 1
            para(doc, "%d. %s" % (li_no, item[1]))
        elif kind == "fig":
            fig_no += 1
            add_figure(doc, item[1], item[2], fig_no)

    doc.save(OUT)
    print("Звіт збережено: %s" % OUT)
    print("Рисунків у звіті: %d" % fig_no)


if __name__ == "__main__":
    main()
